from pathlib import Path

from docx import Document

from brd_agent.extraction.extractors.docx.docx_extractor import extract_docx


def test_extract_docx_returns_heading_paragraph_and_table(tmp_path: Path):
    source = tmp_path / "requirements.docx"

    document = Document()

    document.add_heading("Requirements", level=1)
    document.add_paragraph(
        "The system shall support document uploads."
    )

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Priority"
    table.cell(1, 0).text = "Uploads"
    table.cell(1, 1).text = "High"

    document.save(source)

    result = extract_docx(str(source))

    assert result["document"]["file_type"] == "docx"

    assert result["elements"][0]["type"] == "heading"
    assert result["elements"][1]["type"] == "paragraph"

    table_element = result["elements"][2]

    assert table_element["type"] == "table"
    assert table_element["content"]["headers"] == [
        "Name",
        "Priority",
    ]
    assert table_element["content"]["rows"] == [
        ["Uploads", "High"],
    ]