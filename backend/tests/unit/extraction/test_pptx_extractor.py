from pathlib import Path

from pptx import Presentation

from brd_agent.extraction.extractors.pptx.pptx_extractor import extract_pptx


def test_extract_pptx_returns_slide_text(tmp_path: Path):
    source = tmp_path / "requirements.pptx"

    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = "Requirements"
    slide.placeholders[1].text = (
        "The system shall support document uploads."
    )

    presentation.save(source)

    result = extract_pptx(str(source))

    assert result["document"]["file_type"] == "pptx"
    assert result["document"]["slide_count"] == 1

    assert result["elements"][0]["type"] == "heading"
    assert result["elements"][0]["content"] == "Requirements"

    assert result["elements"][1]["type"] == "paragraph"
    assert "document uploads" in result["elements"][1]["content"]

    assert result["elements"][0]["location"]["slide"] == 1
    assert "bbox" in result["elements"][0]["location"]